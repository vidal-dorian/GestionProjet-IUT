import base64
from datetime import datetime
from io import BytesIO

from docx import Document
from PIL import Image

from app import models
from app.report_content import parse_bullets
from tests.conftest import TestingSessionLocal
from tests.helpers import add_approved_member


def create_test_project(client, name="Projet CR", email="alice@test.local"):
    client.headers["X-Dev-Email"] = email
    return client.post("/api/projects", json={"name": name}).json()


def create_sprint(client, project_id, name="Sprint 3", start="2026-03-06", end="2026-03-23"):
    return client.post(
        f"/api/projects/{project_id}/sprints", json={"name": name, "start_date": start, "end_date": end}
    ).json()["sprint"]


def account_id(client, email):
    previous = client.headers.get("X-Dev-Email")
    client.headers["X-Dev-Email"] = email
    value = client.get("/api/me").json()["id"]
    client.headers["X-Dev-Email"] = previous
    return value


def add_issue(project_id, number, title, labels, state="open"):
    db = TestingSessionLocal()
    try:
        db.add(
            models.GithubIssue(
                project_id=project_id,
                number=number,
                title=title,
                state=state,
                labels_raw=",".join(labels),
                url=f"https://github.com/o/r/issues/{number}",
                synced_at=datetime.utcnow(),
                story_points=3,
                closed_at=datetime(2026, 3, 10) if state == "closed" else None,
            )
        )
        db.commit()
    finally:
        db.close()


def docx_text(response) -> str:
    document = Document(BytesIO(response.content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def save(http, project_id, report, **content_changes):
    return http.put(
        f"/api/projects/{project_id}/reports/{report['id']}",
        json={
            "sprint_id": report["sprint_id"],
            "meeting_date": report["meeting_date"],
            "content": {**report["content"], **content_changes},
            "version": report["version"],
        },
    )


def test_reports_require_membership(client):
    project = create_test_project(client)
    client.headers["X-Dev-Email"] = "intrus@test.local"
    assert client.get(f"/api/projects/{project['id']}/reports").status_code == 403


def test_create_daily_defaults_participants_to_sprint_team(client):
    project = create_test_project(client)
    add_approved_member(project["id"], "bob@test.local")
    add_approved_member(project["id"], "carol@test.local")
    sprint = create_sprint(client, project["id"])
    roles = {r["name"]: r["id"] for r in client.get(f"/api/projects/{project['id']}/roles").json()}
    alice, bob = account_id(client, "alice@test.local"), account_id(client, "bob@test.local")
    client.put(
        f"/api/projects/{project['id']}/sprints/{sprint['id']}/role-assignments",
        json={"assignments": [{"role_id": roles["Développeur"], "account_id": alice}, {"role_id": roles["Product Owner"], "account_id": bob}]},
    )

    response = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "daily", "sprint_id": sprint["id"], "meeting_date": "2026-03-10"},
    )
    assert response.status_code == 201
    report = response.json()
    assert sorted(report["content"]["participant_ids"]) == sorted([alice, bob])
    assert len(report["team"]) == 3
    assert report["can_delete"] is True


def test_only_one_daily_per_date(client):
    project = create_test_project(client)
    payload = {"type": "daily", "meeting_date": "2026-03-10"}
    assert client.post(f"/api/projects/{project['id']}/reports", json=payload).status_code == 201
    assert client.post(f"/api/projects/{project['id']}/reports", json=payload).status_code == 409


def test_sprint_ceremonies_require_a_sprint(client):
    project = create_test_project(client)
    response = client.post(
        f"/api/projects/{project['id']}/reports", json={"type": "sprint_review", "meeting_date": "2026-03-23"}
    )
    assert response.status_code == 422


def test_review_prefills_closed_user_stories_and_planning_objectives(client):
    project = create_test_project(client)
    sprint = create_sprint(client, project["id"])
    add_issue(project["id"], 5, "Demander la modification", ["Sprint 3"], state="closed")
    add_issue(project["id"], 3, "Demande champ facultatif", ["Sprint 3"], state="closed")
    add_issue(project["id"], 7, "Encore ouverte", ["Sprint 3"], state="open")
    add_issue(project["id"], 9, "Autre sprint", ["Sprint 2"], state="closed")

    planning = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_planning", "sprint_id": sprint["id"], "meeting_date": "2026-03-06"},
    ).json()
    assert [row["reference"] for row in planning["content"]["user_stories"]] == ["#3", "#5", "#7"]
    assert save(client, project["id"], planning, topics="Terminer l'API\n  Gérer les demandes", client="Mme Servieres").status_code == 200

    review = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_review", "sprint_id": sprint["id"], "meeting_date": "2026-03-23"},
    ).json()
    assert review["content"]["user_stories"] == [
        {"reference": "#3", "name": "Demande champ facultatif"},
        {"reference": "#5", "name": "Demander la modification"},
    ]
    assert review["content"]["objectives"] == "Terminer l'API\n  Gérer les demandes"
    assert review["content"]["client"] == "Mme Servieres"


def test_suggested_user_stories_reflect_latest_github_state(client):
    project = create_test_project(client)
    sprint = create_sprint(client, project["id"])
    review = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_review", "sprint_id": sprint["id"], "meeting_date": "2026-03-23"},
    ).json()
    assert review["content"]["user_stories"] == []

    add_issue(project["id"], 11, "Supprimer une demande", ["Sprint 3"], state="closed")
    response = client.get(f"/api/projects/{project['id']}/reports/{review['id']}/suggested-user-stories")
    assert response.status_code == 200
    assert response.json()["user_stories"] == [{"reference": "#11", "name": "Supprimer une demande"}]


def test_update_rejects_stale_version(client):
    project = create_test_project(client)
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()

    first = save(client, project["id"], report, scrum_master_notes="RAS")
    assert first.status_code == 200
    assert first.json()["version"] == 2

    stale = save(client, project["id"], report, scrum_master_notes="Écrase tout")
    assert stale.status_code == 409


def test_update_rejects_non_member_participant(client):
    project = create_test_project(client)
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()
    outsider = account_id(client, "intrus@test.local")
    assert save(client, project["id"], report, participant_ids=[outsider]).status_code == 422


def test_update_rejects_invalid_content(client):
    project = create_test_project(client)
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()
    assert save(client, project["id"], report, balance="nulle-part").status_code == 422


def test_each_member_fills_their_own_daily_entry(client):
    project = create_test_project(client)
    add_approved_member(project["id"], "bob@test.local")
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()
    bob = account_id(client, "bob@test.local")

    client.headers["X-Dev-Email"] = "bob@test.local"
    response = client.put(
        f"/api/projects/{project['id']}/reports/{report['id']}/daily-entries/{bob}",
        json={"done": "Installation des VM", "todo": "Export de VM", "blockers": ""},
    )
    assert response.status_code == 200
    entries = response.json()["daily_entries"]
    assert entries[0]["done"] == "Installation des VM"

    # Une seconde écriture met à jour la même ligne au lieu d'en créer une autre.
    response = client.put(
        f"/api/projects/{project['id']}/reports/{report['id']}/daily-entries/{bob}",
        json={"done": "Installation des VM", "todo": "Rédaction livrables", "blockers": ""},
    )
    assert [e["todo"] for e in response.json()["daily_entries"]] == ["Rédaction livrables"]


def test_daily_entry_for_non_member_is_rejected(client):
    project = create_test_project(client)
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()
    outsider = account_id(client, "intrus@test.local")
    response = client.put(
        f"/api/projects/{project['id']}/reports/{report['id']}/daily-entries/{outsider}", json={"done": "x"}
    )
    assert response.status_code == 422


def test_delete_restricted_to_author_or_owner(client):
    project = create_test_project(client)
    add_approved_member(project["id"], "bob@test.local")
    add_approved_member(project["id"], "carol@test.local")
    client.headers["X-Dev-Email"] = "bob@test.local"
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()

    client.headers["X-Dev-Email"] = "carol@test.local"
    assert client.delete(f"/api/projects/{project['id']}/reports/{report['id']}").status_code == 403

    client.headers["X-Dev-Email"] = "alice@test.local"
    assert client.delete(f"/api/projects/{project['id']}/reports/{report['id']}").status_code == 204
    assert client.get(f"/api/projects/{project['id']}/reports").json() == []


def test_daily_export_matches_template(client):
    project = create_test_project(client)
    sprint = create_sprint(client, project["id"], name="Sprint 4", start="2026-03-24", end="2026-04-03")
    client.put("/api/me", json={"display_name": "VIDAL Dorian"})
    roles = {r["name"]: r["id"] for r in client.get(f"/api/projects/{project['id']}/roles").json()}
    alice = account_id(client, "alice@test.local")
    client.put(
        f"/api/projects/{project['id']}/sprints/{sprint['id']}/role-assignments",
        json={"assignments": [{"role_id": roles["Développeur"], "account_id": alice}]},
    )
    report = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "daily", "sprint_id": sprint["id"], "meeting_date": "2026-03-30"},
    ).json()
    client.put(
        f"/api/projects/{project['id']}/reports/{report['id']}/daily-entries/{alice}",
        json={"done": "Rédaction du document technique", "todo": "Préparation des livrables", "blockers": ""},
    )

    response = client.get(f"/api/projects/{project['id']}/reports/{report['id']}/export")
    assert response.status_code == 200
    assert 'filename="Daily_-_30_03_2026.docx"' in response.headers["content-disposition"]
    text = docx_text(response)
    assert "Daily - 30/03/2026" in text
    assert "Sprint n° : 4" in text
    assert "Personnes présentes : VIDAL Dorian" in text
    assert "VIDAL Dorian / Développeur" in text
    assert "Ce que j’ai fait depuis la dernière daily : Rédaction du document technique" in text
    assert "Ce qui me bloque : /" in text
    assert "Notes du Scrum Master" in text


def test_review_export_contains_user_story_and_signature_tables(client):
    project = create_test_project(client)
    client.put(f"/api/projects/{project['id']}/document-settings", json={"footer": "SAE S4 - Développement"})
    sprint = create_sprint(client, project["id"])
    add_issue(project["id"], 3, "Demande pour renseigner un champ facultatif", ["Sprint 3"], state="closed")
    review = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_review", "sprint_id": sprint["id"], "meeting_date": "2026-03-23"},
    ).json()
    save(client, project["id"], review, client="Servieres Corinne", objectives="Terminer l'API\n  Gérer les demandes")

    response = client.get(f"/api/projects/{project['id']}/reports/{review['id']}/export")
    assert response.status_code == 200
    assert 'filename="Review_de_Sprint_3.docx"' in response.headers["content-disposition"]
    text = docx_text(response)
    assert "Réunion de Review" in text
    assert "Itération N°3" in text
    assert "#3 | Demande pour renseigner un champ facultatif" in text
    assert "Client | Approbation" in text
    assert "Servieres Corinne" in text
    footer = Document(BytesIO(response.content)).sections[0].footer.paragraphs[0].text
    assert footer.startswith("SAE S4 - Développement")


def test_retrospective_export_includes_burndown(client):
    project = create_test_project(client)
    sprint = create_sprint(client, project["id"])
    add_issue(project["id"], 3, "US fermée", ["Sprint 3"], state="closed")
    add_issue(project["id"], 4, "US ouverte", ["Sprint 3"])
    retro = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "retrospective", "sprint_id": sprint["id"], "meeting_date": "2026-03-23"},
    ).json()
    save(client, project["id"], retro, went_well="Bonne communication", burndown_comment="Dans cette itération...")

    response = client.get(f"/api/projects/{project['id']}/reports/{retro['id']}/export")
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    assert len(document.inline_shapes) == 1
    text = docx_text(response)
    assert "Rétrospective de Sprint" in text
    assert "Bonne communication" in text
    assert "Dans cette itération..." in text


def test_planning_export_computes_capacity(client):
    project = create_test_project(client)
    add_approved_member(project["id"], "bob@test.local")
    sprint = create_sprint(client, project["id"])
    planning = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_planning", "sprint_id": sprint["id"], "meeting_date": "2026-03-06"},
    ).json()
    save(client, project["id"], planning, hours_per_member=12)

    text = docx_text(client.get(f"/api/projects/{project['id']}/reports/{planning['id']}/export"))
    assert "12 heures de travail sont planifiées pour chaque membre de l’équipe" in text
    assert "Capacité totale : 24 heures" in text


def test_display_name_update(client):
    client.headers["X-Dev-Email"] = "alice@test.local"
    response = client.put("/api/me", json={"display_name": "  ALINDÉ Yoan "})
    assert response.status_code == 200
    assert response.json()["display_name"] == "ALINDÉ Yoan"
    assert client.put("/api/me", json={"display_name": ""}).json()["display_name"] is None


def _png_base64() -> str:
    buffer = BytesIO()
    Image.new("RGB", (40, 20), "red").save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode()


def test_logo_upload_and_export(client):
    project = create_test_project(client)
    response = client.put(
        f"/api/projects/{project['id']}/document-settings/logos/left", json={"data_base64": _png_base64()}
    )
    assert response.status_code == 200
    assert response.json()["logos"] == ["left"]
    logo = client.get(f"/api/projects/{project['id']}/document-settings/logos/left")
    assert logo.headers["content-type"] == "image/png"

    sprint = create_sprint(client, project["id"])
    planning = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_planning", "sprint_id": sprint["id"], "meeting_date": "2026-03-06"},
    ).json()
    exported = client.get(f"/api/projects/{project['id']}/reports/{planning['id']}/export")
    assert len(Document(BytesIO(exported.content)).inline_shapes) == 1

    assert client.delete(f"/api/projects/{project['id']}/document-settings/logos/left").status_code == 204
    assert client.get(f"/api/projects/{project['id']}/document-settings").json()["logos"] == []


def test_logo_upload_rejects_non_images_and_non_owners(client):
    project = create_test_project(client)
    url = f"/api/projects/{project['id']}/document-settings/logos/left"
    not_an_image = base64.b64encode(b"<svg></svg>").decode()
    assert client.put(url, json={"data_base64": not_an_image}).status_code == 422

    add_approved_member(project["id"], "bob@test.local")
    client.headers["X-Dev-Email"] = "bob@test.local"
    assert client.put(url, json={"data_base64": _png_base64()}).status_code == 403


def test_parse_bullets_handles_indentation_and_manual_markers():
    assert parse_bullets("- Terminer l'API\n  - Champs facultatifs\n\n\tGérer les demandes\n• Autre") == [
        (0, "Terminer l'API"),
        (1, "Champs facultatifs"),
        (1, "Gérer les demandes"),
        (0, "Autre"),
    ]


def test_footer_defaults_to_project_name(client):
    project = create_test_project(client, name="Application de covoiturage")
    sprint = create_sprint(client, project["id"])
    review = client.post(
        f"/api/projects/{project['id']}/reports",
        json={"type": "sprint_review", "sprint_id": sprint["id"], "meeting_date": "2026-03-23"},
    ).json()

    response = client.get(f"/api/projects/{project['id']}/reports/{review['id']}/export")
    footer = Document(BytesIO(response.content)).sections[0].footer.paragraphs[0].text
    assert footer.startswith("Application de covoiturage")


def test_daily_balance_point_can_be_placed_freely(client):
    project = create_test_project(client)
    report = client.post(f"/api/projects/{project['id']}/reports", json={"type": "daily", "meeting_date": "2026-03-10"}).json()

    saved = save(client, project["id"], report, balance_x=320.5, balance_y=290)
    assert saved.status_code == 200
    assert (saved.json()["content"]["balance_x"], saved.json()["content"]["balance_y"]) == (320.5, 290)
    assert client.get(f"/api/projects/{project['id']}/reports/{report['id']}/export").status_code == 200

    outside = save(client, project["id"], saved.json(), balance_x=5000, balance_y=290)
    assert outside.status_code == 422


def test_balance_diagram_draws_dot_at_given_point():
    from app.docx_charts import project_balance_diagram

    image = Image.open(project_balance_diagram("all", (150, 600)))
    assert image.size == (1000, 820)
    red, green, _ = image.getpixel((150, 600))
    assert red > 200 and green < 60
    # Plus de point à la position par défaut (centre du diagramme).
    assert image.getpixel((500, 387)) == (255, 255, 255)
