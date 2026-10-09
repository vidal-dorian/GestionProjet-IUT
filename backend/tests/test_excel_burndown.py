import re
import zipfile
from datetime import date

from openpyxl import load_workbook

from app import excel_export, models


def _export():
    sprint = models.Sprint(
        id=1, project_id=1, name="Sprint 0", start_date=date(2026, 9, 1), end_date=date(2026, 9, 14)
    )
    data = {
        "total_points": 10,
        "unestimated_issue_count": 0,
        "ideal": [
            {"date": date(2026, 9, 1), "remaining_points": 10},
            {"date": date(2026, 9, 14), "remaining_points": 0},
        ],
        "actual": [
            {"date": date(2026, 9, 1), "remaining_points": 10},
            {"date": date(2026, 9, 5), "remaining_points": 7},
            {"date": date(2026, 9, 10), "remaining_points": 7},
        ],
    }
    return excel_export.build_burndown_export(sprint, data)


def test_burndown_chart_references_exactly_the_data_rows():
    buffer = _export()
    sheet = load_workbook(buffer).active
    rows = list(sheet.iter_rows(values_only=True))
    header_index = rows.index(("Date", "Idéal (SP restants)", "Réel (SP restants)"))
    first, last = header_index + 2, len(rows)  # numéros de ligne Excel (base 1)

    chart_xml = zipfile.ZipFile(buffer).read("xl/charts/chart1.xml").decode()
    refs = re.findall(r"<f>'Burndown'!\$([A-C])\$(\d+):\$[A-C]\$(\d+)</f>", chart_xml)

    # Ni la ligne d'en-tête ni la ligne vide au-dessus ne sont tracées, et la
    # dernière date (fin du sprint) n'est pas perdue.
    expected_range = (str(first), str(last))
    assert sorted(set(refs)) == [(col, *expected_range) for col in "ABC"]
    assert "<scatterChart>" in chart_xml
    assert '<smooth val="1"/>' not in chart_xml
    assert chart_xml.count('<delete val="0"/>') == 2  # axes visibles


def test_actual_line_stops_at_the_last_known_point():
    sheet = load_workbook(_export()).active
    last_row = list(sheet.iter_rows(values_only=True))[-1]
    assert last_row[0].date() == date(2026, 9, 14)
    assert last_row[1] == 0
    assert last_row[2] is None
