"""Export Word (.docx) des comptes-rendus de cérémonies Scrum.

La mise en page reproduit les modèles utilisés jusqu'ici par l'équipe : le
daily façon Google Docs (titre centré, informations générales, trois
questions par personne, notes du Scrum Master et diagramme Temps / Qualité /
Respect du cahier des charges) ; la planification, la review et la
rétrospective avec une page de titre (logos, "Itération N°x"), un pied de
page et un tableau des US.
"""

import re
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from app import docx_charts
from app.text import strip_control_chars
from app.report_content import (
    DailyContent,
    RetrospectiveContent,
    SprintPlanningContent,
    SprintReviewContent,
    UserStoryRow,
    parse_bullets,
)

SERIF = "Times New Roman"
SANS = "Arial"
EMPTY = "/"


@dataclass
class Participant:
    account_id: int
    label: str
    roles: list[str] = field(default_factory=list)


@dataclass
class DailyAnswers:
    done: str = ""
    todo: str = ""
    blockers: str = ""


@dataclass
class ReportExportData:
    type: str
    meeting_date: date
    sprint_name: str | None
    sprint_start: date | None
    sprint_end: date | None
    content: object
    participants: list[Participant]
    daily_answers: dict[int, DailyAnswers] = field(default_factory=dict)
    footer: str = ""
    logos: dict[str, bytes] = field(default_factory=dict)
    burndown: dict | None = None


def iteration_number(sprint_name: str | None) -> str:
    """"Sprint 3" → "3" ; un nom sans chiffre est repris tel quel."""
    if not sprint_name:
        return EMPTY
    match = re.search(r"\d+", sprint_name)
    return match.group(0) if match else sprint_name


def join_names(names: list[str]) -> str:
    if not names:
        return EMPTY
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + " et " + names[-1]


def _format_number(value: float) -> str:
    return f"{value:g}".replace(".", ",")


# --- Briques de mise en page -------------------------------------------------


def _new_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.54))

    normal = document.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)
    # Sans rFonts/eastAsia, Word garde la police du thème pour certains caractères.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)

    for style_name, size in (("Heading 1", 18), ("Heading 2", 16)):
        style = document.styles[style_name]
        style.font.name = SERIF
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.rPr.rFonts
        rfonts.set(qn("w:eastAsia"), SERIF)
        # Les titres du modèle par défaut suivent la police du thème
        # (asciiTheme), qui l'emporte sur la police explicite.
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attribute), None)
        style.paragraph_format.space_before = Pt(18)
        style.paragraph_format.space_after = Pt(6)
    for style_name in ("List Bullet", "List Bullet 2"):
        document.styles[style_name].paragraph_format.space_after = Pt(2)
    return document


def _run(paragraph, text: str, *, bold: bool = False, size: float | None = None, font: str | None = None):
    # Filet de sécurité : un caractère de contrôle ferait échouer tout l'export.
    run = paragraph.add_run(strip_control_chars(text))
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if font is not None:
        run.font.name = font
        run.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    return run


def _blank(document, count: int = 1) -> None:
    for _ in range(count):
        document.add_paragraph()


def _bullets(document, text: str, *, font: str | None = None, size: float = 11, empty: str | None = EMPTY) -> None:
    items = parse_bullets(text)
    if not items and empty is not None:
        items = [(0, empty)]
    for level, item in items:
        paragraph = document.add_paragraph(style="List Bullet 2" if level else "List Bullet")
        _run(paragraph, item, size=size, font=font)


def _labelled_line(document, label: str, value: str, *, size: float = 13, font: str = SERIF, indent: float | None = None):
    paragraph = document.add_paragraph()
    if indent is not None:
        paragraph.paragraph_format.left_indent = Cm(indent)
    paragraph.paragraph_format.space_after = Pt(2)
    _run(paragraph, label, bold=True, size=size, font=font)
    _run(paragraph, f" : {value}", size=size, font=font)
    return paragraph


def _add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    """Insère un champ Word (PAGE, NUMPAGES...) recalculé à l'ouverture."""
    for kind, text in (("begin", None), ("instr", instruction), ("separate", None), ("text", placeholder), ("end", None)):
        run = paragraph.add_run()
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = f" {text} "
        elif kind == "text":
            element = OxmlElement("w:t")
            element.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


def _footer(document, text: str) -> None:
    section = document.sections[0]
    paragraph = section.footer.paragraphs[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    tab_stops = paragraph.paragraph_format.tab_stops
    # Le style "Footer" du modèle par défaut a des taquets pour du format
    # Letter (centre à 3,25", droite à 6,5") : on les neutralise.
    for position in (Inches(3.25), Inches(6.5)):
        tab_stops.add_tab_stop(position, WD_TAB_ALIGNMENT.CLEAR)
    tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    _run(paragraph, f"{text}\t")
    _add_field(paragraph, "PAGE")
    _run(paragraph, "/")
    _add_field(paragraph, "NUMPAGES")
    for run in paragraph.runs:
        run.font.name = SERIF
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def _remove_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _logos(document, logos: dict[str, bytes]) -> None:
    if not logos:
        return
    table = document.add_table(rows=1, cols=2)
    _remove_borders(table)
    for index, position in enumerate(("left", "right")):
        data = logos.get(position)
        paragraph = table.rows[0].cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if position == "left" else WD_ALIGN_PARAGRAPH.RIGHT
        if data:
            paragraph.add_run().add_picture(BytesIO(data), height=Cm(2))


def _title_page(document, data: ReportExportData, title: str) -> None:
    _logos(document, data.logos)
    _blank(document, 3)
    for line in (title, f"Itération N°{iteration_number(data.sprint_name)}"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _run(paragraph, line, bold=True, size=23, font=SERIF)


def _participant_lines(document, data: ReportExportData, client: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    _run(paragraph, "Participants", bold=True, size=13, font=SERIF)
    _run(paragraph, " :", size=13, font=SERIF)
    names = join_names([p.label for p in data.participants])
    for label, value in (("MOE", names), ("Client", client.strip() or EMPTY)):
        line = document.add_paragraph()
        line.paragraph_format.left_indent = Cm(1.27)
        line.paragraph_format.space_after = Pt(2)
        _run(line, f"{label} : {value}", size=13, font=SERIF)


def _page_break(document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _set_cell(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(paragraph, text, bold=bold, size=11, font=SERIF)


def _set_column_widths(table, widths) -> None:
    # Word lit la largeur des cellules, LibreOffice celle de la grille : on
    # renseigne les deux.
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _user_story_table(document, rows: list[UserStoryRow]) -> None:
    rows = [row for row in rows if row.reference.strip() or row.name.strip()]
    if not rows:
        paragraph = document.add_paragraph()
        _run(paragraph, "Aucune user story.", size=11, font=SERIF)
        return
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell(table.rows[0].cells[0], "Référence", bold=True, center=True)
    _set_cell(table.rows[0].cells[1], "Nom", bold=True, center=True)
    for row in rows:
        cells = table.add_row().cells
        _set_cell(cells[0], row.reference.strip(), center=True)
        _set_cell(cells[1], row.name.strip())
    _set_column_widths(table, (Cm(2.5), Cm(13.5)))


def _signature_table(document, data: ReportExportData, client: str) -> None:
    heading = document.add_heading("Signature des parties prenantes :", level=2)
    heading.paragraph_format.space_after = Pt(12)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def add(left: str, right: str = "", *, bold: bool = False) -> None:
        row = table.add_row()
        row.height = Cm(0.9)
        _set_cell(row.cells[0], left, bold=bold)
        _set_cell(row.cells[1], right, bold=bold)

    add("Membres de l'équipe", "Présence", bold=True)
    for participant in data.participants:
        add(participant.label)
    add("")
    add("Client", "Approbation", bold=True)
    add(client.strip())
    _set_column_widths(table, (Cm(8), Cm(8)))


# --- Documents ---------------------------------------------------------------


def _build_daily(document, data: ReportExportData) -> None:
    content: DailyContent = data.content
    date_text = data.meeting_date.strftime("%d/%m/%Y")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(24)
    _run(title, f"Daily - {date_text}", size=26)
    _blank(document)

    _run(document.add_paragraph(), "Informations générales :", size=15)
    for label, value in (
        ("Date", date_text),
        ("Sprint n°", iteration_number(data.sprint_name)),
        ("Personnes présentes", ", ".join(p.label for p in data.participants) or EMPTY),
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        _run(paragraph, f"{label} : ", bold=True, size=10)
        _run(paragraph, value, size=10)

    questions = (
        ("done", "Ce que j’ai fait depuis la dernière daily"),
        ("todo", "Ce que je vais faire aujourd’hui"),
        ("blockers", "Ce qui me bloque"),
    )
    for participant in data.participants:
        _blank(document)
        heading = document.add_paragraph()
        heading.paragraph_format.left_indent = Cm(0.63)
        heading.paragraph_format.keep_with_next = True
        role = ", ".join(participant.roles)
        _run(heading, f"{participant.label} / {role}" if role else participant.label, bold=True, size=13)

        answers = data.daily_answers.get(participant.account_id, DailyAnswers())
        for number, (key, question) in enumerate(questions, start=1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.9)
            paragraph.paragraph_format.first_line_indent = Cm(-0.63)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(1.9))
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = number < len(questions)
            answer = " ".join(line.strip() for line in getattr(answers, key).splitlines() if line.strip())
            _run(paragraph, f"{number}.\t{question} : {answer or EMPTY}")

    _blank(document, 2)
    heading = document.add_paragraph()
    heading.paragraph_format.left_indent = Cm(0.63)
    _run(heading, "Notes du Scrum Master", bold=True, size=13)
    _bullets(document, content.scrum_master_notes)

    _blank(document, 2)
    point = None
    if content.balance_x is not None and content.balance_y is not None:
        point = (content.balance_x, content.balance_y)
    document.add_picture(docx_charts.project_balance_diagram(content.balance, point), width=Cm(12))


def _build_planning(document, data: ReportExportData) -> None:
    content: SprintPlanningContent = data.content
    _title_page(document, data, "Réunion de Planification")
    _blank(document, 10)
    _labelled_line(document, "Date de la réunion", data.meeting_date.strftime("%d/%m/%Y"))
    _participant_lines(document, data, content.client)
    _page_break(document)

    document.add_heading("Sujet de la réunion", level=1)
    _bullets(document, content.topics, font=SERIF)

    if content.hours_per_member is not None:
        document.add_heading("Estimation du temps de travail pour le sprint", level=2)
        capacity = content.hours_per_member * len(data.participants)
        paragraph = document.add_paragraph()
        hours = _format_number(content.hours_per_member)
        _run(paragraph, f"{hours} heure{'s' if content.hours_per_member > 1 else ''} de travail sont planifiées pour chaque membre de l’équipe")
        _run(paragraph, "\n")
        _run(paragraph, f"Capacité totale : {_format_number(capacity)} heures", bold=True)

    document.add_heading(
        f"Proposition de la MOE des items de backlog à réaliser dans l’itération {iteration_number(data.sprint_name)}",
        level=2,
    )
    _user_story_table(document, content.user_stories)
    _blank(document)
    _signature_table(document, data, content.client)


def _build_review(document, data: ReportExportData) -> None:
    content: SprintReviewContent = data.content
    _title_page(document, data, "Réunion de Review")
    _blank(document, 10)
    _labelled_line(document, "Date de la réunion", data.meeting_date.strftime("%d/%m/%Y"))
    _participant_lines(document, data, content.client)
    _page_break(document)

    document.add_heading("Objectifs du sprint :", level=1)
    _bullets(document, content.objectives, font=SERIF)
    document.add_heading("User Stories réalisées :", level=1)
    _user_story_table(document, content.user_stories)
    _blank(document)
    _signature_table(document, data, content.client)


def _build_retrospective(document, data: ReportExportData) -> None:
    content: RetrospectiveContent = data.content
    _title_page(document, data, "Rétrospective de Sprint")
    _blank(document, 2)
    _labelled_line(document, "Date de la rétrospective", data.meeting_date.strftime("%d/%m/%Y"))
    _labelled_line(document, "Participants", join_names([p.label for p in data.participants]))
    _labelled_line(document, "Scrum Master", content.scrum_master.strip() or EMPTY)
    _labelled_line(document, "Product Owner", content.product_owner.strip() or EMPTY)
    _blank(document, 2)

    for label, text in (
        ("Ce qui a bien fonctionné :", content.went_well),
        ("Ce qui est à améliorer :", content.to_improve),
        ("Actions d'amélioration :", content.actions),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.keep_with_next = True
        _run(paragraph, label, bold=True, size=13, font=SANS)
        _bullets(document, text, font=SANS)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(24)
    paragraph.paragraph_format.keep_with_next = True
    _run(paragraph, "Notes du Product Owner :", bold=True, size=13, font=SANS)
    notes = [line.strip() for line in content.product_owner_notes.splitlines() if line.strip()]
    for line in notes or [EMPTY]:
        _run(document.add_paragraph(), line, size=11, font=SANS)

    if content.include_burndown:
        document.add_heading("Burndown Chart", level=1)
        if data.burndown and data.burndown["matched_issue_count"] > 0 and data.sprint_start and data.sprint_end:
            title = f"Burndown chart itération {iteration_number(data.sprint_name)}"
            chart = docx_charts.burndown_chart(title, data.burndown, data.sprint_start, data.sprint_end)
            document.add_picture(chart, width=Cm(16))
        else:
            _run(
                document.add_paragraph(),
                f"Aucune US GitHub labellisée « {data.sprint_name or ''} » : burndown indisponible.",
                size=11,
                font=SANS,
            )
        for line in [line.strip() for line in content.burndown_comment.splitlines() if line.strip()]:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _run(paragraph, line, size=11, font=SANS)


_BUILDERS = {
    "daily": _build_daily,
    "sprint_planning": _build_planning,
    "sprint_review": _build_review,
    "retrospective": _build_retrospective,
}


def build_report_docx(data: ReportExportData) -> BytesIO:
    document = _new_document()
    _BUILDERS[data.type](document, data)
    if data.type != "daily":
        _footer(document, data.footer.strip())
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
